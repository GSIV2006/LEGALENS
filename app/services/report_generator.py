"""
Report Generator Service

Generates PDF and DOCX compliance reports using ReportLab and python-docx.

PDF includes:
- Legal Metrology Compliance Assessment header
- Inspection ID, product details, date
- Inspector info
- Overall status
- All declaration checks with PASS/FAIL/MANUAL REVIEW
- Violation summary
- Detected text/evidence with confidence
- Legal rule references
- Notes
- Ruleset version
- Image/evidence references where possible
"""
import io
import os
from datetime import datetime
from typing import Dict, List, Optional, Any

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm, inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Image as RLImage,
    PageBreak,
    KeepTogether,
)
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from app.schemas.compliance import ComplianceAnalysis, CheckStatus
from app.schemas.inspection import InspectionResponse
from app.schemas.product import ProductResponse
from app.schemas.auth import UserResponse


class ReportGenerator:
    """
    Service for generating compliance reports in PDF and DOCX formats.
    """

    def __init__(self):
        """Initialize report generator."""
        self.styles = getSampleStyleSheet()
        self._setup_custom_styles()

    def _setup_custom_styles(self):
        """Setup custom paragraph styles."""
        self.styles.add(ParagraphStyle(
            name='ReportTitle',
            parent=self.styles['Title'],
            fontSize=18,
            spaceAfter=12,
            alignment=TA_CENTER,
            textColor=colors.HexColor('#1a5276'),
        ))
        self.styles.add(ParagraphStyle(
            name='SectionHeader',
            parent=self.styles['Heading2'],
            fontSize=14,
            spaceBefore=12,
            spaceAfter=6,
            textColor=colors.HexColor('#2c3e50'),
        ))
        self.styles.add(ParagraphStyle(
            name='SubSection',
            parent=self.styles['Heading3'],
            fontSize=12,
            spaceBefore=8,
            spaceAfter=4,
            textColor=colors.HexColor('#34495e'),
        ))
        self.styles.add(ParagraphStyle(
            name='StatusPass',
            parent=self.styles['Normal'],
            textColor=colors.HexColor('#27ae60'),
            fontSize=10,
        ))
        self.styles.add(ParagraphStyle(
            name='StatusFail',
            parent=self.styles['Normal'],
            textColor=colors.HexColor('#c0392b'),
            fontSize=10,
        ))
        self.styles.add(ParagraphStyle(
            name='StatusReview',
            parent=self.styles['Normal'],
            textColor=colors.HexColor('#f39c12'),
            fontSize=10,
        ))
        self.styles.add(ParagraphStyle(
            name='StatusNA',
            parent=self.styles['Normal'],
            textColor=colors.HexColor('#7f8c8d'),
            fontSize=10,
        ))

    def generate_pdf(
        self,
        inspection: InspectionResponse,
        product: Optional[ProductResponse],
        inspector: Optional[UserResponse],
        compliance_analysis: ComplianceAnalysis,
        ocr_data: Optional[List[Dict]] = None,
        images: Optional[List[Dict]] = None,
        output_path: Optional[str] = None,
    ) -> bytes:
        """
        Generate PDF compliance report.

        Args:
            inspection: Inspection details
            product: Product details (optional)
            inspector: Inspector details (optional)
            compliance_analysis: Complete compliance analysis
            ocr_data: OCR extracted text data (optional)
            images: Image evidence list (optional)
            output_path: Optional file path to save (if None, returns bytes)

        Returns:
            PDF file as bytes
        """
        buffer = io.BytesIO()

        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=20*mm,
            leftMargin=20*mm,
            topMargin=20*mm,
            bottomMargin=20*mm,
        )

        story = []

        # Title
        story.append(Paragraph("LEGAL METROLOGY COMPLIANCE ASSESSMENT", self.styles['ReportTitle']))
        story.append(Spacer(1, 6*mm))
        story.append(Paragraph(
            "Packaged Commodities (Legal Metrology) Rules, 2011",
            ParagraphStyle('subtitle', parent=self.styles['Normal'],
                          fontSize=11, alignment=TA_CENTER, textColor=colors.grey)
        ))
        story.append(Spacer(1, 10*mm))

        # Header Information
        story.append(self._create_header_table(inspection, product, inspector))
        story.append(Spacer(1, 8*mm))

        # Overall Status
        story.append(self._create_status_section(compliance_analysis))
        story.append(Spacer(1, 8*mm))

        # Compliance Checks
        story.append(Paragraph("COMPLIANCE DECLARATION CHECKS", self.styles['SectionHeader']))
        story.append(self._create_checks_table(compliance_analysis))
        story.append(Spacer(1, 8*mm))

        # Violation Summary
        if compliance_analysis.failed_checks > 0:
            story.append(self._create_violation_summary(compliance_analysis))

        # Evidence / Detected Text
        if ocr_data and len(ocr_data) > 0:
            story.append(Paragraph("DETECTED TEXT / EVIDENCE", self.styles['SectionHeader']))
            story.append(self._create_evidence_section(ocr_data))

        # Ruleset Information
        story.append(Paragraph("RULESET INFORMATION", self.styles['SectionHeader']))
        story.append(self._create_ruleset_table(compliance_analysis, inspection))
        story.append(Spacer(1, 8*mm))

        # Notes
        if inspection.notes:
            story.append(Paragraph("INSPECTOR NOTES", self.styles['SectionHeader']))
            story.append(Paragraph(inspection.notes, self.styles['Normal']))
            story.append(Spacer(1, 6*mm))

        # Footer
        story.append(Spacer(1, 10*mm))
        story.append(Paragraph(
            f"Report generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')} UTC",
            ParagraphStyle('footer', parent=self.styles['Normal'],
                          fontSize=9, textColor=colors.grey, alignment=TA_CENTER)
        ))
        story.append(Paragraph(
            "This is a prototype compliance assessment. "
            "Legal verification by qualified professional is recommended.",
            ParagraphStyle('disclaimer', parent=self.styles['Normal'],
                          fontSize=8, textColor=colors.red, alignment=TA_CENTER)
        ))

        # Build PDF
        doc.build(story)

        pdf_bytes = buffer.getvalue()
        buffer.close()

        # Save to file if path provided
        if output_path:
            with open(output_path, "wb") as f:
                f.write(pdf_bytes)

        return pdf_bytes

    def _create_header_table(
        self,
        inspection: InspectionResponse,
        product: Optional[ProductResponse],
        inspector: Optional[UserResponse],
    ) -> Table:
        """Create header information table."""
        data = [
            ["Inspection ID:", str(inspection.id)],
            ["Status:", inspection.status],
            ["Ruleset Version:", inspection.ruleset_version],
        ]

        if product:
            data.append(["Product:", product.product_name])
            if product.brand:
                data.append(["Brand:", product.brand])
            if product.category:
                data.append(["Category:", product.category])
            if product.manufacturer:
                data.append(["Manufacturer:", product.manufacturer])

        if inspector:
            data.append(["Inspector:", inspector.full_name or inspector.username])
            data.append(["Inspector Email:", inspector.email])

        data.append(["Inspection Date:", inspection.created_at.strftime('%Y-%m-%d %H:%M') if inspection.created_at else "N/A"])
        data.append(["Last Updated:", inspection.updated_at.strftime('%Y-%m-%d %H:%M') if inspection.updated_at else "N/A"])

        # Add image count if available
        if inspection.images_count and inspection.images_count > 0:
            data.append(["Images Uploaded:", str(inspection.images_count)])

        # Add OCR data count
        if inspection.ocr_data_count and inspection.ocr_data_count > 0:
            data.append(["OCR Text Items:", str(inspection.ocr_data_count)])

        # Add compliance checks count
        if inspection.checks_count and inspection.checks_count > 0:
            data.append(["Compliance Checks:", str(inspection.checks_count)])

        table = Table(data, colWidths=[120, 400])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#ecf0f1')),
            ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#2c3e50')),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#bdc3c7')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))

        return table

    def _create_status_section(self, analysis: ComplianceAnalysis) -> Table:
        """Create overall status section."""
        status_color = {
            CheckStatus.COMPLIANT: colors.HexColor('#27ae60'),
            CheckStatus.NON_COMPLIANT: colors.HexColor('#c0392b'),
            CheckStatus.MANUAL_REVIEW: colors.HexColor('#f39c12'),
            CheckStatus.NOT_APPLICABLE: colors.HexColor('#7f8c8d'),
            CheckStatus.NOT_VERIFIED: colors.HexColor('#95a5a6'),
        }.get(analysis.overall_status, colors.black)

        data = [
            [
                Paragraph(f"OVERALL STATUS: {analysis.overall_status.upper()}", 
                         ParagraphStyle('status', parent=self.styles['Normal'],
                                       fontSize=14, textColor=status_color,
                                       alignment=TA_CENTER)),
            ],
            [
                Paragraph(f"Compliance Percentage: {analysis.compliance_percentage}%", 
                         self.styles['Normal']),
            ],
        ]

        table = Table(data, colWidths=[520])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f8f9fa')),
            ('BOX', (0, 0), (-1, -1), 1, status_color),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ]))

        return table

    def _create_checks_table(self, analysis: ComplianceAnalysis) -> Table:
        """Create compliance checks table."""
        header = ["Rule Code", "Field", "Status", "Message", "Confidence", "Evidence"]
        data = [header]

        for check in analysis.checks:
            status_style = {
                CheckStatus.PASS: 'StatusPass',
                CheckStatus.FAIL: 'StatusFail',
                CheckStatus.MANUAL_REVIEW: 'StatusReview',
                CheckStatus.NOT_APPLICABLE: 'StatusNA',
                CheckStatus.NOT_VERIFIED: 'StatusNA',
            }.get(check.status, 'Normal')

            evidence = check.evidence or "-"
            if len(evidence) > 60:
                evidence = evidence[:57] + "..."

            data.append([
                check.rule_code,
                check.field,
                Paragraph(check.status.upper(), self.styles[status_style]),
                check.message or "-",
                f"{check.confidence:.0%}" if check.confidence else "-",
                evidence,
            ])

        table = Table(data, colWidths=[70, 90, 80, 150, 50, 80])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2c3e50')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#bdc3c7')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
        ]))

        return table

    def _create_violation_summary(self, analysis: ComplianceAnalysis) -> Table:
        """Create violation summary section."""
        data = [
            [Paragraph("VIOLATION SUMMARY", self.styles['SubSection'])],
            [Paragraph(
                f"Total Violations: {analysis.failed_checks} | "
                f"Manual Review Required: {analysis.manual_review_checks}",
                self.styles['Normal']
            )],
        ]

        for check in analysis.checks:
            if check.status == CheckStatus.FAIL:
                data.append([
                    Paragraph(f"<b>{check.rule_code}</b>: {check.message}", self.styles['Normal'])
                ])
                if check.evidence and len(check.evidence) > 80:
                    data.append([
                        Paragraph(f"Evidence: {check.evidence[:77]}...", 
                                 ParagraphStyle('evidence', parent=self.styles['Normal'],
                                              fontSize=9, textColor=colors.grey))
                    ])

        table = Table(data, colWidths=[520])
        table.setStyle(TableStyle([
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#c0392b')),
            ('BACKGROUND', (0, 0), (0, 0), colors.HexColor('#fdedec')),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('LEFTPADDING', (0, 0), (-1, -1), 10),
        ]))

        return table

    def _create_evidence_section(self, ocr_data: List[Dict]) -> Table:
        """Create detected text/evidence section."""
        data = [["Source Image", "Confidence", "Detected Text"]]

        for item in ocr_data[:30]:  # Limit to 30 items
            text = item.get('text', '')[:80]
            if len(item.get('text', '')) > 80:
                text += "..."

            confidence = f"{item.get('confidence', 0):.0%}" if item.get('confidence') else "-"
            source = item.get('source_image', 'unknown') or 'unknown'

            data.append([source, confidence, text])

        table = Table(data, colWidths=[80, 60, 380])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#34495e')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#bdc3c7')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8f9fa')]),
        ]))

        return table

    def _create_ruleset_table(
        self,
        analysis: ComplianceAnalysis,
        inspection: InspectionResponse,
    ) -> Table:
        """Create ruleset information table."""
        data = [
            ["Ruleset Version:", analysis.ruleset_version],
            ["Analysis Date:", analysis.analyzed_at.strftime('%Y-%m-%d %H:%M:%S') if analysis.analyzed_at else "N/A"],
            ["Total Rules Checked:", str(analysis.total_checks)],
            ["Passed:", str(analysis.passed_checks)],
            ["Failed:", str(analysis.failed_checks)],
            ["Manual Review:", str(analysis.manual_review_checks)],
            ["Not Applicable:", str(analysis.not_applicable_checks)],
        ]

        table = Table(data, colWidths=[120, 400])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#ecf0f1')),
            ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#2c3e50')),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#bdc3c7')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ]))

        return table

    def generate_docx(
        self,
        inspection: InspectionResponse,
        product: Optional[ProductResponse],
        inspector: Optional[UserResponse],
        compliance_analysis: ComplianceAnalysis,
        ocr_data: Optional[List[Dict]] = None,
        images: Optional[List[Dict]] = None,
        output_path: Optional[str] = None,
    ) -> bytes:
        """
        Generate DOCX compliance report.

        Args:
            inspection: Inspection details
            product: Product details (optional)
            inspector: Inspector details (optional)
            compliance_analysis: Complete compliance analysis
            ocr_data: OCR extracted text data (optional)
            images: Image evidence list (optional)
            output_path: Optional file path to save (if None, returns bytes)

        Returns:
            DOCX file as bytes
        """
        doc = Document()

        # Set document styles
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)

        # Title
        title = doc.add_heading('LEGAL METROLOGY COMPLIANCE ASSESSMENT', level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Subtitle
        subtitle = doc.add_paragraph(
            'Packaged Commodities (Legal Metrology) Rules, 2011'
        )
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.runs[0].italic = True

        doc.add_paragraph()  # Spacer

        # Header Info
        doc.add_heading('Inspection Details', level=1)
        self._add_docx_info_table(doc, inspection, product, inspector)

        doc.add_paragraph()

        # Overall Status
        doc.add_heading('Overall Compliance Status', level=1)
        status_para = doc.add_paragraph()
        status_run = status_para.add_run(analysis.overall_status.upper())
        status_run.bold = True
        status_run.font.size = Pt(14)

        status_colors = {
            CheckStatus.COMPLIANT: RGBColor(39, 174, 96),
            CheckStatus.NON_COMPLIANT: RGBColor(192, 57, 43),
            CheckStatus.MANUAL_REVIEW: RGBColor(243, 156, 18),
            CheckStatus.NOT_VERIFIED: RGBColor(149, 165, 166),
        }
        if analysis.overall_status in status_colors:
            status_run.font.color.rgb = status_colors[analysis.overall_status]

        doc.add_paragraph(f"Compliance Percentage: {analysis.compliance_percentage}%")
        doc.add_paragraph()

        # Compliance Checks
        doc.add_heading('Compliance Declaration Checks', level=1)
        self._add_docx_checks_table(doc, compliance_analysis)
        doc.add_paragraph()

        # Violation Summary
        if analysis.failed_checks > 0:
            doc.add_heading('Violation Summary', level=1)
            for check in compliance_analysis.checks:
                if check.status == CheckStatus.FAIL:
                    p = doc.add_paragraph()
                    p.add_run(f"{check.rule_code}: ").bold = True
                    p.add_run(check.message)
                    if check.evidence:
                        doc.add_paragraph(f"Evidence: {check.evidence}")
            doc.add_paragraph()

        # Detected Text / Evidence
        if ocr_data and len(ocr_data) > 0:
            doc.add_heading('Detected Text / Evidence', level=1)
            self._add_docx_evidence_table(doc, ocr_data)
            doc.add_paragraph()

        # Ruleset Information
        doc.add_heading('Ruleset Information', level=1)
        self._add_docx_ruleset_table(doc, compliance_analysis, inspection)
        doc.add_paragraph()

        # Notes
        if inspection.notes:
            doc.add_heading('Inspector Notes', level=1)
            doc.add_paragraph(inspection.notes)
            doc.add_paragraph()

        # Footer
        doc.add_paragraph()
        footer_para = doc.add_paragraph()
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_para.add_run(
            f"Report generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')} UTC"
        ).italic = True

        disclaimer = doc.add_paragraph()
        disclaimer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        disclaimer.add_run(
            "This is a prototype compliance assessment. "
            "Legal verification by qualified professional is recommended."
        ).italic = True

        # Save to bytes
        buffer = io.BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.close()

        # Save to file if path provided
        if output_path:
            with open(output_path, "wb") as f:
                f.write(docx_bytes)

        return docx_bytes

    def _add_docx_info_table(
        self,
        doc: Document,
        inspection: InspectionResponse,
        product: Optional[ProductResponse],
        inspector: Optional[UserResponse],
    ):
        """Add inspection info table to DOCX."""
        table = doc.add_table(rows=0, cols=2)
        table.style = 'Light Grid Accent 1'

        def add_row(label, value):
            row = table.add_row().cells
            row[0].text = label
            row[1].text = str(value) if value else "-"
            # Make label bold
            for paragraph in row[0].paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        add_row("Inspection ID", str(inspection.id))
        add_row("Status", inspection.status)
        add_row("Ruleset Version", inspection.ruleset_version)
        add_row("Inspection Date", 
                inspection.created_at.strftime('%Y-%m-%d %H:%M') if inspection.created_at else "-")
        add_row("Last Updated", 
                inspection.updated_at.strftime('%Y-%m-%d %H:%M') if inspection.updated_at else "-")

        if product:
            add_row("Product Name", product.product_name)
            if product.brand:
                add_row("Brand", product.brand)
            if product.category:
                add_row("Category", product.category)
            if product.manufacturer:
                add_row("Manufacturer", product.manufacturer)

        if inspector:
            add_row("Inspector", inspector.full_name or inspector.username)
            add_row("Inspector Email", inspector.email)

        if inspection.images_count:
            add_row("Images Uploaded", str(inspection.images_count))
        if inspection.ocr_data_count:
            add_row("OCR Text Items", str(inspection.ocr_data_count))
        if inspection.checks_count:
            add_row("Compliance Checks", str(inspection.checks_count))

    def _add_docx_checks_table(
        self,
        doc: Document,
        analysis: ComplianceAnalysis,
    ):
        """Add compliance checks table to DOCX."""
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Light Grid Accent 1'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header
        header_cells = table.rows[0].cells
        headers = ["Rule Code", "Field", "Status", "Message", "Evidence"]
        for i, header in enumerate(headers):
            header_cells[i].text = header
            for paragraph in header_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.size = Pt(9)

        # Data rows
        status_colors = {
            CheckStatus.PASS: RGBColor(39, 174, 96),
            CheckStatus.FAIL: RGBColor(192, 57, 43),
            CheckStatus.MANUAL_REVIEW: RGBColor(243, 156, 18),
            CheckStatus.NOT_APPLICABLE: RGBColor(127, 140, 141),
            CheckStatus.NOT_VERIFIED: RGBColor(149, 165, 166),
        }

        for check in analysis.checks:
            row = table.add_row().cells
            row[0].text = check.rule_code
            row[1].text = check.field
            row[2].text = check.status
            row[3].text = check.message or "-"
            row[4].text = (check.evidence or "-")[:100]

            # Color the status cell
            for paragraph in row[2].paragraphs:
                for run in paragraph.runs:
                    if check.status in status_colors:
                        run.font.color.rgb = status_colors[check.status]
                    run.bold = True

    def _add_docx_evidence_table(
        self,
        doc: Document,
        ocr_data: List[Dict],
    ):
        """Add OCR evidence table to DOCX."""
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Light Grid Accent 1'

        # Header
        header_cells = table.rows[0].cells
        headers = ["Source Image", "Confidence", "Detected Text"]
        for i, header in enumerate(headers):
            header_cells[i].text = header
            for paragraph in header_cells[i].paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        # Data
        for item in ocr_data[:50]:
            row = table.add_row().cells
            row[0].text = item.get('source_image', 'unknown') or 'unknown'
            row[1].text = f"{item.get('confidence', 0):.0%}" if item.get('confidence') else "-"
            row[2].text = (item.get('text', '') or '')[:150]

    def _add_docx_ruleset_table(
        self,
        doc: Document,
        analysis: ComplianceAnalysis,
        inspection: InspectionResponse,
    ):
        """Add ruleset info table to DOCX."""
        table = doc.add_table(rows=0, cols=2)
        table.style = 'Light Grid Accent 1'

        def add_row(label, value):
            row = table.add_row().cells
            row[0].text = label
            row[1].text = str(value) if value else "-"
            for paragraph in row[0].paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        add_row("Ruleset Version", analysis.ruleset_version)
        add_row("Analysis Date", 
                analysis.analyzed_at.strftime('%Y-%m-%d %H:%M:%S') if analysis.analyzed_at else "-")
        add_row("Total Rules Checked", str(analysis.total_checks))
        add_row("Passed", str(analysis.passed_checks))
        add_row("Failed", str(analysis.failed_checks))
        add_row("Manual Review Required", str(analysis.manual_review_checks))
        add_row("Not Applicable", str(analysis.not_applicable_checks))


# Global instance
report_generator = ReportGenerator()
